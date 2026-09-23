"""Renders a ContractSummary into a professionally formatted Word document, styled to
match the visual language of pdf_report.py / report.pdf.html.jinja (navy headings) so all
output formats read as one product family."""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor

from contract_reviewer.schemas import ContractSummary

_NAVY = RGBColor(0x1A, 0x2B, 0x4C)
_GREY = RGBColor(0x55, 0x55, 0x55)


def _style_heading(paragraph, size: int | None = None) -> None:
    for run in paragraph.runs:
        run.font.color.rgb = _NAVY
        if size:
            run.font.size = Pt(size)


def _add_heading(document: Document, text: str, level: int):
    heading = document.add_heading(text, level=level)
    _style_heading(heading)
    return heading


def _add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def _add_two_col_table(document: Document, header: tuple[str, str], rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    header_cells[0].text = header[0]
    header_cells[1].text = header[1]
    for cell in header_cells:
        cell.paragraphs[0].runs[0].bold = True
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value or "not stated"


def render_docx(summary: ContractSummary, context: dict) -> bytes:
    document = Document()

    title = document.add_heading("Contract Summary", level=0)
    _style_heading(title, size=26)

    meta = document.add_paragraph()
    meta.add_run(f"Source: {context.get('source_filename', '')}\n").italic = True
    meta.add_run(f"Generated: {context.get('generated_at', '')}\n").italic = True
    if context.get("client_role"):
        meta.add_run(f"Prepared for: {context['client_role']}\n").italic = True
    if context.get("source_sha256"):
        small = meta.add_run(f"Source SHA-256: {context['source_sha256']}")
        small.italic = True
        small.font.size = Pt(7)
        small.font.color.rgb = _GREY

    _add_heading(document, "Snapshot", 1)
    _add_two_col_table(
        document,
        ("Field", "Detail"),
        [
            ("Contract type", summary.contract_type),
            ("Parties", "; ".join(f"{p.name} ({p.role})" for p in summary.parties) or "not stated"),
            ("Effective date", summary.effective_date),
            ("Term / duration", summary.term_duration),
            ("Governing law", summary.governing_law),
            ("Dispute resolution", summary.dispute_resolution),
            ("Key monetary value", summary.key_monetary_value),
        ],
    )

    _add_heading(document, "Subject Matter and Purpose", 1)
    document.add_paragraph(summary.subject_matter_summary)

    _add_heading(document, "Key Obligations by Party", 1)
    for party_obligations in summary.obligations_by_party:
        _add_heading(document, party_obligations.party, 2)
        _add_bullets(document, party_obligations.obligations)

    _add_heading(document, "Payment Terms", 1)
    document.add_paragraph(summary.payment_terms)

    if summary.key_dates:
        _add_heading(document, "Key Dates and Milestones", 1)
        _add_two_col_table(
            document, ("Date", "Description"), [(d.date, d.description) for d in summary.key_dates]
        )

    _add_heading(document, "Termination and Exit", 1)
    document.add_paragraph(summary.termination_and_exit)

    _add_heading(document, "Liability and Risk-Allocation Snapshot", 1)
    document.add_paragraph(summary.liability_snapshot)

    if summary.ip_confidentiality_data:
        _add_heading(document, "IP, Confidentiality and Data", 1)
        document.add_paragraph(summary.ip_confidentiality_data)

    if summary.notable_or_unusual_terms:
        _add_heading(document, "Notable or Unusual Terms", 1)
        _add_bullets(document, summary.notable_or_unusual_terms)

    if summary.missing_information:
        _add_heading(document, "Missing Information", 1)
        _add_bullets(document, summary.missing_information)

    fidelity_issues = context.get("fidelity_issues") or []
    if fidelity_issues:
        _add_heading(document, "Data Fidelity Notices", 1)
        _add_bullets(document, fidelity_issues)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
