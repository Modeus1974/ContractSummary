"""Extracts text from a contract Word document (.docx).

Returns the same shape as pdf_extract.extract_pdf() so intake_node can treat both
document types identically. The source file is treated as immutable evidence: this
module only reads it.
"""
from __future__ import annotations

from pathlib import Path

import docx


def extract_docx(path: Path) -> dict:
    """Returns {ok, text, pages, warnings, scanned_pages}. Never raises on a bad/corrupt
    .docx; extraction failure is reported in the result so the caller can route to the
    Incomplete path. A .docx has no real page concept, so `pages` is a single logical
    entry (page_no=1) rather than an approximation."""
    result = {
        "ok": False,
        "text": "",
        "pages": [],
        "warnings": [],
        "scanned_pages": [],
    }
    try:
        document = docx.Document(path)
    except Exception as exc:  # corrupt / not a valid .docx
        result["warnings"].append(f"Could not open Word document: {exc}")
        return result

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    table_lines = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                table_lines.append(" | ".join(cells))

    text_parts = paragraphs + table_lines
    combined_text = "\n".join(text_parts)

    result["pages"] = [{"page_no": 1, "text": combined_text}]
    result["text"] = f"[Page 1]\n{combined_text}" if combined_text.strip() else ""
    result["ok"] = bool(combined_text.strip())
    if not result["ok"]:
        result["warnings"].append("No extractable text found anywhere in the document.")
    return result
